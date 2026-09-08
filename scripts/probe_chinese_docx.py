"""Independent in-memory DOCX integration probe; no LLM and no user documents."""
import base64
import io
import json
from pathlib import Path
import sys
from xml.etree import ElementTree as ET
import zipfile

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / '.vendor/docxengine/src'))
from docx import Document
from docxengine import (Session, ToolError, docx_open, docx_comment, docx_replace,
                        docx_revision, docx_validate, export_bytes, paragraph_anchor)

NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}

def open_bytes(data):
    session = Session()
    opened = docx_open(session, bytes=base64.b64encode(data).decode('ascii'))
    return session, opened['doc_id']

def revision_signature(session, doc_id, author):
    records = docx_revision(session, doc_id=doc_id, op='list')['revisions']
    return [(r['id'], r['type'], r['text'], r['author'], r['date'])
            for r in records if r['author'] == author]

def package_parts(data):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}

def main():
    doc = Document()
    doc.add_paragraph('本科毕业论文')
    doc.add_paragraph('本研究使用卷积神经网络（CNN）进行分类。')
    paragraph = doc.add_paragraph()
    paragraph.add_run('本研究')
    paragraph.add_run('非常').bold = True
    paragraph.add_run('非常')
    paragraph.add_run('有效。')
    doc.add_comment(paragraph.runs, text='避免主观评价，请给出实验依据。', author='老师甲', initials='甲')
    doc.sections[0].header.paragraphs[0].text = '毕业论文技术验证'
    table = doc.add_table(rows=2, cols=2)
    table.cell(0,0).text = '模型'
    table.cell(0,1).text = '准确率'
    table.cell(1,0).text = '示例模型'
    table.cell(1,1).text = '示例值'
    buffer = io.BytesIO()
    doc.save(buffer)
    original = buffer.getvalue()
    session, doc_id = open_bytes(original)
    assert docx_comment(session, doc_id=doc_id, op='list')['comments'][0]['text'] == '避免主观评价，请给出实验依据。'

    docx_replace(session, doc_id=doc_id,
                 anchor=paragraph_anchor(2, '本研究使用卷积神经网络（CNN）进行分类。'),
                 old='分类', new='图像分类', track_changes=True, author='老师甲')
    reviewed = export_bytes(session, doc_id=doc_id)
    session, doc_id = open_bytes(reviewed)
    teacher_before = revision_signature(session, doc_id, '老师甲')
    assert len(teacher_before) == 2
    original_comments = docx_comment(session, doc_id=doc_id, op='list')['comments']
    anchor = paragraph_anchor(3, '本研究非常非常有效。')
    result = docx_replace(session, doc_id=doc_id, anchor=anchor,
                          old='非常非常', new='较为', track_changes=True, author='审改助手')
    assert result['n_replaced'] == 1
    assert result['new_anchor'] == paragraph_anchor(3, '本研究较为有效。')
    docx_comment(session, doc_id=doc_id, op='add', anchor=result['new_anchor'],
                 text='示例修改仅验证修订功能；正式建议仍需实验依据。', author='审改助手')
    assert revision_signature(session, doc_id, '老师甲') == teacher_before
    assert len(revision_signature(session, doc_id, '审改助手')) == 2
    current_comments = docx_comment(session, doc_id=doc_id, op='list')['comments']
    assert len(current_comments) == 2
    assert current_comments[0]['text'] == original_comments[0]['text']
    assert current_comments[0]['author'] == '老师甲'
    validation = docx_validate(session, doc_id=doc_id)
    assert validation['valid'], validation
    result_bytes = export_bytes(session, doc_id=doc_id)
    before = package_parts(reviewed)
    after = package_parts(result_bytes)
    assert before['word/header1.xml'] == after['word/header1.xml']
    before_root = ET.fromstring(before['word/document.xml'])
    after_root = ET.fromstring(after['word/document.xml'])
    assert ET.tostring(before_root.find('.//w:tbl', NS)) == ET.tostring(after_root.find('.//w:tbl', NS))
    assert after_root.find('.//w:b', NS) is not None
    Document(io.BytesIO(result_bytes))  # independent library parses the saved package

    try:
        docx_comment(session, doc_id=doc_id, op='add', anchor=anchor,
                     text='此定位已过期，应拒绝。', author='审改助手')
    except ToolError as exc:
        assert exc.code == 'anchor_stale', exc
    else:
        raise AssertionError('Stale anchor was accepted')
    rejected_anchor_bytes = export_bytes(session, doc_id=doc_id)
    raw_failure_details = {
        'changed_parts': [name for name, content in package_parts(rejected_anchor_bytes).items()
                          if after.get(name) != content],
        'before_size': len(result_bytes), 'after_size': len(rejected_anchor_bytes)
    }
    # Keep the upstream failure visible. Discard its damaged session and verify a
    # host-side transaction boundary using a fresh, isolated copy of the last good bytes.
    assert raw_failure_details['changed_parts'] == ['word/comments.xml'], raw_failure_details
    staging_session, staging_id = open_bytes(result_bytes)
    committed_bytes = result_bytes
    try:
        docx_comment(staging_session, doc_id=staging_id, op='add', anchor=anchor,
                     text='过期定位测试', author='审改助手')
        committed_bytes = export_bytes(staging_session, doc_id=staging_id)
    except ToolError as exc:
        assert exc.code == 'anchor_stale', exc
    else:
        raise AssertionError('Expected the staged operation to fail')
    assert committed_bytes == result_bytes

    reject_session, reject_id = open_bytes(result_bytes)
    docx_revision(reject_session, doc_id=reject_id, op='reject', filter={'author': '审改助手'})
    assert revision_signature(reject_session, reject_id, '老师甲') == teacher_before
    assert revision_signature(reject_session, reject_id, '审改助手') == []
    rejected = Document(io.BytesIO(export_bytes(reject_session, doc_id=reject_id)))
    assert rejected.paragraphs[2].text == '本研究非常非常有效。'
    report = {
        'probe': 'Chinese DOCX integration, in memory, no model calls',
        'checks_passed': ['read existing teacher comments', 'Chinese cross-run tracked replacement',
                          'retain teacher revisions', 'append separate assistant comment',
                          'retain header bytes and table XML', 'isolated staging retains committed bytes on failure',
                          'independent python-docx parse', 'reject assistant revisions while retaining teacher revisions'],
        'upstream_defect': {'case': 'adding a comment with a stale anchor raises anchor_stale but mutates comments.xml',
                            **raw_failure_details},
        'candidate_status': 'not approved for direct production integration; staging mitigation tested, raw defect remains',
        'engine_validation': validation,
        'not_tested': ['Microsoft Word opening and rendering', 'real thesis layout',
                       'semantic review quality', 'historical issue matching', 'Pi integration']
    }
    print(json.dumps(report, ensure_ascii=False))

if __name__ == '__main__':
    main()
