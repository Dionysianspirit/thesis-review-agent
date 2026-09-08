"""Shared in-memory DOCX builders for tests. No real student papers."""
from __future__ import annotations

import io

from docx import Document


def sample_history_v1() -> bytes:
    doc = Document()
    doc.add_paragraph("本科毕业论文")
    doc.add_paragraph("本研究使用卷积神经网络（CNN）进行分类。")
    paragraph = doc.add_paragraph()
    paragraph.add_run("本研究")
    paragraph.add_run("非常").bold = True
    paragraph.add_run("非常")
    paragraph.add_run("有效。")
    doc.add_comment(paragraph.runs, text="避免主观评价，请给出实验依据。", author="老师甲", initials="甲")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "模型"
    table.cell(0, 1).text = "准确率"
    table.cell(1, 0).text = "示例模型"
    table.cell(1, 1).text = "示例值"
    return _save(doc)


def sample_history_v2() -> bytes:
    doc = Document()
    doc.add_paragraph("本科毕业论文")
    doc.add_paragraph("本研究使用卷积神经网络（CNN）进行图像分类。")
    paragraph = doc.add_paragraph()
    paragraph.add_run("本研究非常非常有效。")
    doc.add_comment(paragraph.runs, text="此句仍缺实验依据。", author="老师甲", initials="甲")
    return _save(doc)


def sample_same_heading_fixed_body(*, history: bool) -> bytes:
    doc = Document()
    heading = doc.add_paragraph("3.1 实验设计")
    if history:
        doc.add_comment(heading.runs, text="此处缺少实验步骤与评价指标。", author="老师甲", initials="甲")
        doc.add_paragraph("本节尚未展开。")
    else:
        doc.add_paragraph("本节给出数据集划分、训练轮次与准确率、F1 两项评价指标。")
    return _save(doc)


def sample_new_draft() -> bytes:
    doc = Document()
    doc.add_paragraph("本科毕业论文")
    doc.add_paragraph("本研究使用卷积神经网络（CNN）进行图像分类。")
    doc.add_paragraph("本研究非常非常有效。")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "模型"
    table.cell(0, 1).text = "准确率"
    table.cell(1, 0).text = "示例模型"
    table.cell(1, 1).text = "示例值"
    doc.add_paragraph("参考文献")
    doc.add_paragraph("[1] 张三. 示例文献. 期刊, 2024.")
    doc.add_paragraph("[3] 李四. 另一篇文献. 期刊, 2025.")
    return _save(doc)


def _save(doc: Document) -> bytes:
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
