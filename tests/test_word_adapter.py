"""Word adapter: comments, tracked edits, and failed writes stay isolated."""
from __future__ import annotations

import io
from xml.etree import ElementTree as ET
import zipfile

import pytest
from docx import Document

from tests.helpers import sample_history_v1
from thesis_review.errors import ReviewError
from thesis_review.word.adapter import WordAdapter


NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


def _parts(data: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def test_extracts_existing_teacher_comment():
    adapter = WordAdapter()
    opened = adapter.open_bytes(sample_history_v1())
    comments = adapter.extract_comments(opened)
    assert comments[0].author == "老师甲"
    assert comments[0].text == "避免主观评价，请给出实验依据。"
    assert "非常非常有效" in comments[0].span


def test_cross_run_chinese_tracked_replace_keeps_teacher_comment():
    adapter = WordAdapter()
    opened = adapter.open_bytes(sample_history_v1())
    paragraphs = adapter.list_paragraphs(opened)
    target = next(item for item in paragraphs if "非常非常有效" in item.text)
    result = adapter.replace_tracked(
        opened,
        anchor=target.anchor,
        old="非常非常",
        new="较为",
        author="审改助手",
    )
    assert result.n_replaced == 1
    assert "较为有效" in result.new_text
    comments = adapter.extract_comments(opened)
    assert comments[0].author == "老师甲"
    exported = adapter.export_bytes(opened)
    Document(io.BytesIO(exported))


def test_stale_anchor_comment_does_not_commit_bytes():
    adapter = WordAdapter()
    opened = adapter.open_bytes(sample_history_v1())
    paragraphs = adapter.list_paragraphs(opened)
    target = next(item for item in paragraphs if "非常非常有效" in item.text)
    stale = target.anchor
    adapter.replace_tracked(
        opened,
        anchor=stale,
        old="非常非常",
        new="较为",
        author="审改助手",
    )
    before = adapter.export_bytes(opened)
    with pytest.raises(ReviewError) as caught:
        adapter.add_comment(
            opened,
            anchor=stale,
            text="此定位已过期，应拒绝。",
            author="审改助手",
        )
    assert caught.value.code == "anchor_stale"
    after = adapter.export_bytes(opened)
    assert after == before
    changed = [name for name, content in _parts(after).items() if _parts(before).get(name) != content]
    assert changed == []


def test_reject_assistant_revisions_restores_sentence():
    adapter = WordAdapter()
    opened = adapter.open_bytes(sample_history_v1())
    paragraphs = adapter.list_paragraphs(opened)
    target = next(item for item in paragraphs if "非常非常有效" in item.text)
    adapter.replace_tracked(
        opened,
        anchor=target.anchor,
        old="非常非常",
        new="较为",
        author="审改助手",
    )
    adapter.reject_revisions(opened, author="审改助手")
    restored = Document(io.BytesIO(adapter.export_bytes(opened)))
    texts = [paragraph.text for paragraph in restored.paragraphs]
    assert any("本研究非常非常有效。" in text for text in texts)


def test_header_and_table_bytes_survive_comment():
    adapter = WordAdapter()
    opened = adapter.open_bytes(sample_history_v1())
    before = adapter.export_bytes(opened)
    paragraphs = adapter.list_paragraphs(opened)
    target = next(item for item in paragraphs if "卷积神经网络" in item.text)
    adapter.add_comment(opened, anchor=target.anchor, text="核对模型名称。", author="审改助手")
    after = adapter.export_bytes(opened)
    before_root = ET.fromstring(_parts(before)["word/document.xml"])
    after_root = ET.fromstring(_parts(after)["word/document.xml"])
    assert ET.tostring(before_root.find(".//w:tbl", NS)) == ET.tostring(after_root.find(".//w:tbl", NS))
