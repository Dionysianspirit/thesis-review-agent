from __future__ import annotations

import io
from pathlib import Path

from docx import Document

OVERCLAIM_CLAIM_QUOTE = "实验结果表明该方法显著提升了分类准确率。"
OVERCLAIM_EVIDENCE_QUOTE = "准确率由 0.81 提高到 0.83。"
SUPPORTED_CLAIM_QUOTE = "实验结果表明准确率达到 0.91，显著高于基线 0.72。"
SUPPORTED_EVIDENCE_QUOTE = "准确率为 0.91，基线模型为 0.72。配对检验 p<0.01。"


def write_demo_drafts(directory: Path) -> dict[str, Path]:
    directory.mkdir(parents=True, exist_ok=True)
    mapping = {
        "v1": _history_v1(),
        "v2": _history_v2(),
        "new": _new_draft(),
    }
    paths: dict[str, Path] = {}
    for name, data in mapping.items():
        path = directory / f"{name}.docx"
        path.write_bytes(data)
        paths[name] = path
    (directory / "overclaim.docx").write_bytes(overclaim_draft())
    return paths


def overclaim_draft() -> bytes:
    """Conclusion overclaims a tiny accuracy bump with no significance test."""
    doc = Document()
    doc.add_paragraph("本科毕业论文")
    doc.add_paragraph("1 摘要")
    doc.add_paragraph("本文研究一种用于图像分类的卷积神经网络方法。")
    doc.add_paragraph("2 方法")
    doc.add_paragraph("本文使用卷积神经网络提取图像特征并完成分类。")
    doc.add_paragraph("3 实验结果")
    doc.add_paragraph(f"在测试集上，{OVERCLAIM_EVIDENCE_QUOTE}未报告显著性检验，也未给出基线对照。")
    doc.add_paragraph("4 结论")
    doc.add_paragraph(OVERCLAIM_CLAIM_QUOTE)
    doc.add_paragraph("参考文献")
    doc.add_paragraph("[1] 张三. 示例文献. 期刊, 2024.")
    return _save(doc)


def supported_claim_draft() -> bytes:
    """Conclusion claim is backed by a baseline and a significance test."""
    doc = Document()
    doc.add_paragraph("本科毕业论文")
    doc.add_paragraph("1 摘要")
    doc.add_paragraph("本文研究一种用于图像分类的卷积神经网络方法。")
    doc.add_paragraph("2 方法")
    doc.add_paragraph("本文使用卷积神经网络提取图像特征并完成分类。")
    doc.add_paragraph("3 实验结果")
    doc.add_paragraph(SUPPORTED_EVIDENCE_QUOTE)
    doc.add_paragraph("4 结论")
    doc.add_paragraph(SUPPORTED_CLAIM_QUOTE)
    doc.add_paragraph("参考文献")
    doc.add_paragraph("[1] 张三. 示例文献. 期刊, 2024.")
    return _save(doc)


def _history_v1() -> bytes:
    doc = Document()
    doc.add_paragraph("本科毕业论文")
    doc.add_paragraph("本研究使用卷积神经网络（CNN）进行分类。")
    paragraph = doc.add_paragraph()
    paragraph.add_run("本研究")
    paragraph.add_run("非常").bold = True
    paragraph.add_run("非常")
    paragraph.add_run("有效。")
    doc.add_comment(paragraph.runs, text="避免主观评价，请给出实验依据。", author="老师甲", initials="甲")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "模型"
    table.cell(0, 1).text = "准确率"
    table.cell(1, 0).text = "示例模型"
    table.cell(1, 1).text = "示例值"
    return _save(doc)


def _history_v2() -> bytes:
    doc = Document()
    doc.add_paragraph("本科毕业论文")
    doc.add_paragraph("本研究使用卷积神经网络（CNN）进行图像分类。")
    paragraph = doc.add_paragraph()
    paragraph.add_run("本研究非常非常有效。")
    doc.add_comment(paragraph.runs, text="此句仍缺实验依据。", author="老师甲", initials="甲")
    return _save(doc)


def _new_draft() -> bytes:
    doc = Document()
    doc.add_paragraph("本科毕业论文")
    doc.add_paragraph("本研究使用卷积神经网络（CNN）进行图像分类。")
    doc.add_paragraph("本研究非常非常有效。")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "模型"
    table.cell(0, 1).text = "准确率"
    table.cell(1, 0).text = "示例模型"
    table.cell(1, 1).text = "示例值"
    doc.add_paragraph("参考文献")
    doc.add_paragraph("[1] 张三. 示例文献. 期刊, 2024.")
    doc.add_paragraph("[3] 李四. 另一篇文献. 期刊, 2025.")
    return _save(doc)


FULL_ABSTRACT_RATE = "准确率为 81%"
FULL_CONCLUSION_RATE = "准确率为 85%"
FULL_OVERCLAIM = "实验结果表明该方法显著提升了分类准确率。"
FULL_SUPPORTED = "在公开数据集上，准确率为 0.91，基线模型为 0.72，配对检验 p<0.01，该对比有统计依据。"
FULL_STATS = "根据国家统计局数据，2023 年相关产业规模已超过十万亿元。"
FULL_METHOD_GAP = "本文使用卷积神经网络完成分类。"
FULL_HISTORY_SPAN = "非常非常有效"


def full_thesis_draft() -> bytes:
    """Synthetic undergraduate thesis: long enough for a first-pass review, with planted issues."""
    doc = Document()
    doc.add_paragraph("大连财经学院本科毕业论文")
    doc.add_paragraph("摘要")
    doc.add_paragraph(
        "本文研究一种用于图像分类的卷积神经网络方法。在测试集上，"
        f"{FULL_ABSTRACT_RATE}。该方法非常非常有效，能够满足毕业设计要求。"
    )
    doc.add_paragraph("关键词：图像分类；卷积神经网络；准确率")
    doc.add_paragraph("Abstract")
    doc.add_paragraph("This paper studies a convolutional neural network for image classification.")
    doc.add_paragraph("1 引言")
    doc.add_paragraph("图像分类是计算机视觉的基础任务。本文提出一种改进卷积网络，并在后文给出实验。")
    doc.add_paragraph("前文将说明数据集划分、训练设置与评价指标，后文应给出对应实验细节。")
    doc.add_paragraph("2 方法")
    doc.add_paragraph(FULL_METHOD_GAP)
    doc.add_paragraph("网络包含卷积层与全连接层。本节未说明学习率、训练轮次、数据增强与划分比例。")
    doc.add_paragraph("3 实验")
    doc.add_paragraph("实验在自建数据集上进行。对照设置仅给出本文模型，未报告基线模型训练细节。")
    doc.add_paragraph("表 1 主要结果")
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "模型"
    table.cell(0, 1).text = "准确率"
    table.cell(1, 0).text = "本文方法"
    table.cell(1, 1).text = "81%"
    table.cell(2, 0).text = "基线"
    table.cell(2, 1).text = "—"
    doc.add_paragraph("4 实验结果")
    doc.add_paragraph(f"在测试集上，{FULL_ABSTRACT_RATE}。未报告显著性检验。")
    doc.add_paragraph(FULL_SUPPORTED)
    doc.add_paragraph("5 结论")
    doc.add_paragraph(f"{FULL_OVERCLAIM}同时，{FULL_CONCLUSION_RATE}。")
    doc.add_paragraph(FULL_STATS)
    doc.add_paragraph("6 参考文献")
    doc.add_paragraph("[1] 张三. 示例文献. 期刊, 2024.")
    doc.add_paragraph("[2] 李四. 另一篇文献. 期刊, 2025.")
    doc.add_paragraph("致谢")
    doc.add_paragraph("感谢指导教师的修改意见。")
    return _save(doc)


def _save(doc: Document) -> bytes:
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
